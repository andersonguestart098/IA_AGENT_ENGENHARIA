# app/ingest/drive_index_pipeline.py
from __future__ import annotations

import io
import os
import re
import json
import hashlib
from datetime import datetime, timezone
from typing import Dict, Any, List, Optional, Tuple

from app.services.drive_store import list_new_files, mark_indexed, mark_error
from app.ingest.qdrant_indexer import (
    get_qdrant,
    ensure_collection,
    upsert_points,
    delete_by_file_id,
)
from app.drive.changes import DriveChangesClient
from app.drive.sheets import SheetsValuesClient
from app.services.sheet_snapshot_store import upsert_sheet_snapshot

QDRANT_COLLECTION = os.getenv("QDRANT_COLLECTION", "drive_rag")
EMBED_MODEL = os.getenv("EMBED_MODEL", "sentence-transformers/all-MiniLM-L6-v2")

# chunking default (texto geral)
CHUNK_SIZE = int(os.getenv("RAG_CHUNK_SIZE", "1200"))
CHUNK_OVERLAP = int(os.getenv("RAG_CHUNK_OVERLAP", "150"))

# chunking planilha (linhas)
SHEET_ROWS_PER_CHUNK = int(os.getenv("SHEET_ROWS_PER_CHUNK", "60"))
SHEET_MAX_COLS = int(os.getenv("SHEET_MAX_COLS", "60"))

_model = None


def _utc_iso() -> str:
    return datetime.now(timezone.utc).isoformat()


def _get_model():
    global _model
    if _model is not None:
        return _model
    from sentence_transformers import SentenceTransformer

    _model = SentenceTransformer(EMBED_MODEL)
    return _model


def _embed_texts(texts: List[str]) -> List[List[float]]:
    model = _get_model()
    return model.encode(texts, normalize_embeddings=True).tolist()


def _sha1(s: str) -> str:
    return hashlib.sha1(s.encode("utf-8", errors="ignore")).hexdigest()


def _norm_ws(s: str) -> str:
    s = s.replace("\x00", " ")
    s = re.sub(r"[ \t]+\n", "\n", s)
    s = re.sub(r"\n{3,}", "\n\n", s)
    s = re.sub(r"[ \t]{2,}", " ", s)
    return s.strip()


def _chunk_text(text: str, chunk_size: int = CHUNK_SIZE, overlap: int = CHUNK_OVERLAP) -> List[str]:
    text = _norm_ws(text)
    if not text:
        return []
    if chunk_size <= 0:
        return [text]

    chunks: List[str] = []
    start = 0
    n = len(text)
    while start < n:
        end = min(start + chunk_size, n)
        chunk = text[start:end].strip()
        if chunk:
            chunks.append(chunk)
        if end == n:
            break
        start = max(0, end - overlap)
    return chunks


# -------------------------
# Drive MIME types
# -------------------------

GOOGLE_DOC = "application/vnd.google-apps.document"
GOOGLE_SHEET = "application/vnd.google-apps.spreadsheet"
GOOGLE_SLIDE = "application/vnd.google-apps.presentation"
GOOGLE_FOLDER = "application/vnd.google-apps.folder"

EXPORT_MAP = {
    GOOGLE_DOC: ("text/plain", "txt"),
    GOOGLE_SLIDE: ("text/plain", "txt"),
    # ⚠️ Sheets NÃO usamos export pra valores (usamos Sheets API)
    # Mantém export xlsx apenas se tu quiser baixar binário por algum motivo:
    # GOOGLE_SHEET: ("application/vnd.openxmlformats-officedocument.spreadsheetml.sheet", "xlsx"),
}

BINARY_SUPPORTED = {
    "application/pdf": "pdf",
    "application/vnd.openxmlformats-officedocument.wordprocessingml.document": "docx",
    "text/plain": "txt",
    "text/markdown": "md",
    "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet": "xlsx",
}


# -------------------------
# Extractors (binários)
# -------------------------

def _extract_text_from_bytes(data: bytes, mime_type: str) -> str:
    # PDF
    if mime_type == "application/pdf":
        try:
            from pypdf import PdfReader
            reader = PdfReader(io.BytesIO(data))
            pages = []
            for p in reader.pages:
                pages.append(p.extract_text() or "")
            return "\n\n".join(pages)
        except Exception:
            return ""

    # DOCX
    if mime_type == "application/vnd.openxmlformats-officedocument.wordprocessingml.document":
        try:
            import docx  # python-docx
            d = docx.Document(io.BytesIO(data))
            return "\n".join([p.text for p in d.paragraphs if p.text])
        except Exception:
            return ""

    # XLSX (fallback para arquivos xlsx uploadados)
    if mime_type == "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet":
        try:
            import openpyxl
            wb = openpyxl.load_workbook(io.BytesIO(data), data_only=True)
            out: List[str] = []
            for sheet in wb.worksheets:
                out.append(f"## Sheet: {sheet.title}")
                for row in sheet.iter_rows(values_only=True):
                    vals = [str(v).strip() for v in row if v is not None and str(v).strip() != ""]
                    if vals:
                        out.append(" | ".join(vals))
                out.append("")
            return "\n".join(out)
        except Exception:
            return ""

    # Texto
    if mime_type in ("text/plain", "text/markdown"):
        try:
            return data.decode("utf-8", errors="ignore")
        except Exception:
            return ""

    return ""


# -------------------------
# Sheets (valores calculados, sem depender de fórmula)
# -------------------------

def _infer_type(v: Any) -> str:
    if v is None:
        return "null"
    if isinstance(v, bool):
        return "bool"
    if isinstance(v, int):
        return "int"
    if isinstance(v, float):
        return "float"
    return "str"


def _rows_to_structured(values: List[List[Any]], *, max_cols: int = SHEET_MAX_COLS) -> Optional[Dict[str, Any]]:
    """
    Converte matriz (Sheets values) em:
    - header: [..]
    - rows: [{col: value}, ...]
    - schema: {col: type}
    """
    if not values or len(values) < 1:
        return None

    header = values[0][:max_cols]
    header = [str(h).strip() if h is not None else "" for h in header]

    # remove colunas vazias no final
    while header and header[-1] == "":
        header.pop()

    if not header:
        return None

    rows_out: List[Dict[str, Any]] = []
    schema: Dict[str, str] = {h: "null" for h in header}

    for row in values[1:]:
        row = row[:len(header)]
        obj: Dict[str, Any] = {}
        for i, col in enumerate(header):
            v = row[i] if i < len(row) else None
            obj[col] = v

            t = _infer_type(v)
            # sobe schema (se algum valor existir, prioriza tipo não-null; se misturar, vira str)
            if schema[col] == "null" and t != "null":
                schema[col] = t
            elif schema[col] != "null" and t != "null" and schema[col] != t:
                schema[col] = "str"

        rows_out.append(obj)

    return {
        "header": header,
        "rows": rows_out,
        "row_count": len(rows_out),
        "schema": schema,
    }


def _sheet_rows_to_chunks_text(
    *,
    sheet_name: str,
    header: List[str],
    rows: List[Dict[str, Any]],
    rows_per_chunk: int = SHEET_ROWS_PER_CHUNK,
) -> List[Tuple[int, int, str]]:
    """
    Retorna lista de chunks: (row_start, row_end, text)
    row_start/end são índices 0-based em relação a rows (dados, sem header).
    """
    chunks: List[Tuple[int, int, str]] = []
    if not rows:
        return chunks

    def fmt(v: Any) -> str:
        if v is None:
            return ""
        # evita notação científica vindo do Sheets: normalmente já vem float “ok”
        if isinstance(v, float):
            s = f"{v:.10f}".rstrip("0").rstrip(".")
            return s
        return str(v).strip()

    for i in range(0, len(rows), rows_per_chunk):
        block = rows[i:i + rows_per_chunk]
        lines: List[str] = []
        lines.append(f"## Sheet: {sheet_name}")
        lines.append("HEADER: " + " | ".join([h for h in header if h]))
        lines.append("ROWS:")
        for j, r in enumerate(block):
            # mostra “linha” compacta com colunas preenchidas
            parts = []
            for col in header:
                val = fmt(r.get(col))
                if val != "":
                    parts.append(f"{col}={val}")
            if parts:
                lines.append(f"L{i + j + 1}: " + " ; ".join(parts))

        text = "\n".join(lines).strip()
        if text:
            chunks.append((i, min(i + rows_per_chunk, len(rows)) - 1, text))

    return chunks


# -------------------------
# Shared payload
# -------------------------

def _build_payload_base(doc: Dict[str, Any]) -> Dict[str, Any]:
    return {
        "source": "gdrive",
        "file_id": doc.get("file_id"),
        "name": doc.get("name"),
        "parent_folder_id": doc.get("parent_folder_id"),
        "parent_folder_name": doc.get("parent_folder_name"),
        "mime_type": doc.get("mime_type"),
        "modified_time": doc.get("modified_time"),
        "size": doc.get("size"),
    }


def _load_sa_info_from_env() -> Dict[str, Any]:
    raw = os.getenv("GOOGLE_SERVICE_ACCOUNT_JSON")
    if not raw:
        raise RuntimeError("GOOGLE_SERVICE_ACCOUNT_JSON não configurado")
    return json.loads(raw)


def _get_drive_client_from_env() -> DriveChangesClient:
    sa_info = _load_sa_info_from_env()
    return DriveChangesClient(service_account_info=sa_info)


def _get_sheets_client_from_env() -> SheetsValuesClient:
    sa_info = _load_sa_info_from_env()
    return SheetsValuesClient(service_account_info=sa_info)


def _download_or_export(drive: DriveChangesClient, file_id: str, mime_type: str) -> Tuple[bytes, str]:
    """
    Retorna (bytes, effective_mime_type)
    - google docs/slides -> export
    - binários -> download
    """
    if mime_type in EXPORT_MAP:
        export_mime, _ext = EXPORT_MAP[mime_type]
        data = drive.export_file_bytes(file_id=file_id, export_mime=export_mime)
        return data, export_mime

    if mime_type in BINARY_SUPPORTED:
        data = drive.download_file_bytes(file_id=file_id)
        return data, mime_type

    data = drive.download_file_bytes(file_id=file_id)
    return data, mime_type


# -------------------------
# MAIN PIPELINE
# -------------------------

async def index_new_drive_files(limit: int = 25) -> Dict[str, Any]:
    docs = await list_new_files(limit=limit)
    if not docs:
        return {"ok": True, "indexed": 0, "skipped": "no_new"}

    drive = _get_drive_client_from_env()
    sheets = _get_sheets_client_from_env()

    qdrant = get_qdrant()

    indexed = 0
    errors = 0
    failed: List[str] = []

    totals = {"files": len(docs), "points": 0, "sheet_rows": 0, "sheet_tabs": 0}

    for doc in docs:
        file_id = doc.get("file_id")
        mime_type = doc.get("mime_type")

        if not file_id or not mime_type:
            await mark_error(file_id or "?", "missing file_id or mime_type")
            errors += 1
            failed.append(file_id or "?")
            continue

        # pula pasta
        if mime_type == GOOGLE_FOLDER:
            await mark_indexed(file_id)
            indexed += 1
            continue

        try:
            # reindex limpo: remove pontos antigos no Qdrant por file_id
            delete_by_file_id(qdrant, QDRANT_COLLECTION, file_id)

            payload_base = _build_payload_base(doc)

            # ==========================================
            # GOOGLE SHEETS (robusto, sem None)
            # ==========================================
            if mime_type == GOOGLE_SHEET:
                sheet_names = sheets.list_sheets(file_id)

                total_file_points = 0

                for sheet_name in sheet_names:
                    values = sheets.get_values(file_id, sheet_name)
                    structured = _rows_to_structured(values)

                    if not structured:
                        continue

                    header = structured["header"]
                    rows = structured["rows"]
                    schema = structured["schema"]

                    totals["sheet_tabs"] += 1
                    totals["sheet_rows"] += len(rows)

                    # snapshot estruturado (determinístico)
                    snapshot_doc = {
                        "file_id": file_id,
                        "file_name": doc.get("name"),
                        "modified_time": doc.get("modified_time"),
                        "parent_folder_id": doc.get("parent_folder_id"),
                        "parent_folder_name": doc.get("parent_folder_name"),
                        "sheet": sheet_name,
                        "header": header,
                        "rows": rows,
                        "row_count": structured["row_count"],
                        "schema": schema,
                        "indexed_at": _utc_iso(),
                    }
                    await upsert_sheet_snapshot(snapshot_doc)

                    # chunks textuais para Qdrant
                    chunks = _sheet_rows_to_chunks_text(
                        sheet_name=sheet_name,
                        header=header,
                        rows=rows,
                        rows_per_chunk=SHEET_ROWS_PER_CHUNK,
                    )

                    if not chunks:
                        continue

                    texts = [t for (_rs, _re, t) in chunks]
                    vectors = _embed_texts(texts)

                    vector_size = len(vectors[0])
                    ensure_collection(qdrant, QDRANT_COLLECTION, vector_size)

                    points: List[Dict[str, Any]] = []
                    for idx, ((row_start, row_end, text), vec) in enumerate(zip(chunks, vectors)):
                        payload = dict(payload_base)
                        payload.update(
                            {
                                "doc_kind": "sheet",
                                "sheet": sheet_name,
                                "row_start": row_start,
                                "row_end": row_end,
                                "chunk_index": idx,
                                "chunk_sha1": _sha1(text),
                                "snapshot_ref": f"{file_id}:{sheet_name}",
                                "text": text,
                            }
                        )
                        points.append(
                            {
                                "id": f"{file_id}:{sheet_name}:{idx}",
                                "vector": vec,
                                "payload": payload,
                            }
                        )

                    upsert_points(qdrant, QDRANT_COLLECTION, points)
                    totals["points"] += len(points)
                    total_file_points += len(points)

                # terminou sheets
                await mark_indexed(file_id)
                indexed += 1
                continue

            # ==========================================
            # OUTROS (PDF/DOCX/TXT/...)
            # ==========================================
            data, effective_mime = _download_or_export(drive, file_id, mime_type)
            text = _extract_text_from_bytes(data, effective_mime)

            if not text or len(text.strip()) < 10:
                await mark_error(file_id, f"no_text_extracted mime={mime_type} effective={effective_mime}")
                errors += 1
                failed.append(file_id)
                continue

            chunks = _chunk_text(text)
            if not chunks:
                await mark_error(file_id, "chunking produced 0 chunks")
                errors += 1
                failed.append(file_id)
                continue

            vectors = _embed_texts(chunks)
            vector_size = len(vectors[0])
            ensure_collection(qdrant, QDRANT_COLLECTION, vector_size)

            points = []
            for i, (chunk, vec) in enumerate(zip(chunks, vectors)):
                payload = dict(payload_base)
                payload.update(
                    {
                        "doc_kind": "text",
                        "chunk_index": i,
                        "chunk_sha1": _sha1(chunk),
                        "text": chunk,
                    }
                )
                points.append({"id": f"{file_id}:{i}", "vector": vec, "payload": payload})

            upsert_points(qdrant, QDRANT_COLLECTION, points)
            totals["points"] += len(points)

            await mark_indexed(file_id)
            indexed += 1

        except Exception as e:
            await mark_error(file_id, f"{type(e).__name__}: {e}")
            errors += 1
            failed.append(file_id)

    return {
        "ok": True,
        "indexed": indexed,
        "errors": errors,
        "total_files": totals["files"],
        "total_points": totals["points"],
        "sheet_tabs": totals["sheet_tabs"],
        "sheet_rows": totals["sheet_rows"],
        "failed": failed[:20],
    }